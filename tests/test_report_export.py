"""REPORT-4I pure tests — export document model, CSV formatter, filename
convention, and format-honesty contract. No database.

Database-backed scope-isolation and end-to-end proofs live in
test_report_export_db.py.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import datetime, timezone

import pytest

from config.reports import get_report
from services import report_export as ex


NOW = datetime(2026, 8, 25, 12, 0, 0, tzinfo=timezone.utc)


@dataclass(frozen=True)
class FakeAlarmRow:
    ou = None
    zone = None
    sector = None
    cnc = None
    feeder_name = None
    transformer: str = "tc1"
    uid: str | None = "29101"
    battery_voltage: float | None = 3.52
    alarm_at: datetime = datetime(2026, 8, 20, 6, 30, tzinfo=timezone.utc)
    temperature: float | None = 85.25
    alarm_label: str = "Battery Alarm"
    firmware_version: str | None = "v7"


def _parse(content: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(content, newline="")))


# ---------------------------------------------------------------------------
# R4-D1 — the honesty contract
# ---------------------------------------------------------------------------


class TestFormatHonesty:
    def test_label_names_development_defaults_and_pending_approval(self):
        label = ex.EXPORT_FORMAT_LABEL.lower()
        assert "development" in label
        assert "csv" in label
        assert "pdf" in label
        assert "xlsx" in label
        assert "pending" in label or "still" in label

    def test_label_never_claims_official_status(self):
        label = ex.EXPORT_FORMAT_LABEL.lower()
        for forbidden in ("required", "official"):
            assert forbidden not in label
        # The label may honestly SAY approval is pending; it must never
        # claim any format IS the approved production mechanism.
        assert "still pending" in label

    def test_label_names_xlsx_as_the_functional_specification_format(self):
        """FS-EXPORT-1: XLSX is not invented — it is what the source's own
        "Export to Excel" UI calls for; the label must say so honestly."""
        assert "excel" in ex.EXPORT_FORMAT_LABEL.lower()

    def test_all_three_formats_are_registered(self):
        """R4-D1/D10/D13: CSV, PDF and native XLSX."""
        assert set(ex.FORMATTERS) == {"csv", "pdf", "xlsx"}

    def test_xlsx_formatter_is_format_xlsx(self):
        assert ex.FORMATTERS["xlsx"] is ex.format_xlsx


# ---------------------------------------------------------------------------
# R4-D9 — which reports are exportable
# ---------------------------------------------------------------------------


class TestExportableReports:
    def test_all_three_real_reports_are_exportable(self):
        assert set(ex.EXPORTABLE_REPORTS) == {
            "installed_rtls", "rtl_alarms_30d", "max_temperature",
        }

    def test_unknown_report_key_has_no_download_path(self):
        with pytest.raises(ValueError):
            ex.render_export(
                "not-a-real-report",
                ex.ExportDocument("not-a-real-report", "t", "fleet", NOW,
                                  ("A",), ()),
            )

    def test_unknown_export_format_is_refused(self):
        doc = ex.installed_rtls_document([], scope_label="fleet", now=NOW)
        with pytest.raises(ValueError):
            ex.render_export("installed_rtls", doc, "ods")


# ---------------------------------------------------------------------------
# R4-D5 — column order and header fidelity
# ---------------------------------------------------------------------------


class TestHeaders:
    @pytest.mark.parametrize(
        "report_key,builder",
        [
            ("rtl_alarms_30d", ex.rtl_alarms_document),
            ("installed_rtls", ex.installed_rtls_document),
        ],
    )
    def test_headers_match_the_client_contract_exactly(self, report_key, builder):
        doc = builder([], scope_label="fleet", now=NOW)   # header-only document
        assert doc.headers == get_report(report_key).columns

    def test_max_temperature_headers_match_the_client_contract_exactly(self):
        doc = ex.max_temperature_document(
            [], scope_label="fleet", period_text="2026-08-08 to 2026-09-07 UTC",
            now=NOW,
        )
        assert doc.headers == get_report("max_temperature").columns

    def test_header_only_document_for_zero_rows(self):
        """R4-D5: a valid zero-row report exports headers only."""
        content = ex.format_csv(ex.rtl_alarms_document([], scope_label="fleet", now=NOW))
        parsed = _parse(content)
        assert len(parsed) == 1
        assert parsed[0][0] == "OU"

    def test_max_temperature_zero_rows_exports_headers_only_no_preamble(self):
        """R4-D5 + R4-D11: a plain header-only file, no metadata preamble —
        Maximum Temperature's CSV is rectangular just like the other two."""
        doc = ex.max_temperature_document(
            [], scope_label="fleet", period_text="2026-08-08 to 2026-09-07 UTC",
            now=NOW,
        )
        parsed = _parse(ex.format_csv(doc))
        assert len(parsed) == 1
        assert parsed[0][0] == "OU"
        assert len(parsed[0]) == len(get_report("max_temperature").columns)


# ---------------------------------------------------------------------------
# R4-D4 — domain-native cells, UTF-8, escaping
# ---------------------------------------------------------------------------


class TestCellSemantics:
    def _alarm_csv(self, row: FakeAlarmRow) -> list[list[str]]:
        return _parse(ex.format_csv(ex.rtl_alarms_document([row], scope_label="fleet", now=NOW)))

    def test_none_becomes_empty_never_the_ui_placeholder(self):
        """R2-D2/R3-D5/R4-D4: '—' must never leak into exported bytes."""
        row = FakeAlarmRow(battery_voltage=None, temperature=None,
                           firmware_version=None, uid=None)
        record = self._alarm_csv(row)[1]
        assert "—" not in ex.format_csv(ex.rtl_alarms_document([row], scope_label="fleet", now=NOW))
        assert [record[7], record[9], record[11], record[6]] == ["", "", "", ""]

    def test_timestamps_are_iso8601_utc(self):
        record = self._alarm_csv(FakeAlarmRow())[1]
        assert record[8] == "2026-08-20T06:30:00Z"

    def test_naive_datetime_is_interpreted_as_utc(self):
        naive = FakeAlarmRow(alarm_at=datetime(2026, 8, 20, 6, 30))
        record = self._alarm_csv(naive)[1]
        assert record[8] == "2026-08-20T06:30:00Z"

    def test_numerics_render_raw_not_decorated(self):
        record = self._alarm_csv(FakeAlarmRow())[1]
        assert record[7] == "3.52"
        assert record[9] == "85.25"

    def test_taxonomy_cells_are_empty(self):
        record = self._alarm_csv(FakeAlarmRow())[1]
        assert record[0:5] == ["", "", "", "", ""]

    def test_commas_quotes_newlines_and_unicode_round_trip(self):
        tricky = 'Alarm, with "quotes" and\nnewline — °C ünicode'
        doc = ex.ExportDocument(
            report_key="rtl_alarms_30d", title="t", scope_label="fleet",
            generated_at_utc=NOW, headers=("Alarm",), rows=[(tricky,)],
        )
        parsed = _parse(ex.format_csv(doc))
        assert parsed[1][0] == tricky

    def test_output_is_plain_utf8_no_bom(self):
        """Review refinement: neutral UTF-8, no BOM, unless client/tooling
        evidence later demands utf-8-sig."""
        doc = ex.rtl_alarms_document([FakeAlarmRow()], scope_label="fleet", now=NOW)
        raw = ex.format_csv(doc).encode("utf-8")
        assert not raw.startswith(b"\xef\xbb\xbf")


# ---------------------------------------------------------------------------
# R4-D6 — deterministic development filename convention
# ---------------------------------------------------------------------------


class TestFilename:
    def test_shape_is_deterministic(self):
        name = ex.export_filename("rtl_alarms_30d", "fleet", NOW)
        assert name == "rtl_alarms_30d_fleet_20260825T120000Z.csv"

    def test_hostile_scope_characters_are_sanitized(self):
        name = ex.export_filename(
            "installed_rtls",
            'plant/"Koeberg",\n unit <1>',
            NOW,
        )
        assert set(name) <= set(
            "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789._-"
        )
        assert name.endswith("_20260825T120000Z.csv")

    def test_long_scope_labels_are_bounded(self):
        name = ex.export_filename("installed_rtls", "x" * 500, NOW)
        assert len(name) <= 120

    def test_render_export_returns_mime_and_filename(self):
        doc = ex.rtl_alarms_document([FakeAlarmRow()], scope_label="fleet", now=NOW)
        content, mime, filename = ex.render_export("rtl_alarms_30d", doc)
        assert mime == "text/csv"
        assert filename == "rtl_alarms_30d_fleet_20260825T120000Z.csv"
        assert "Battery Alarm" in content

    def test_max_temperature_filename_carries_the_period(self):
        """R4-D11: only Maximum Temperature's filename gains a period
        segment — the other two reports' filenames are unchanged (proven
        above by test_shape_is_deterministic)."""
        name = ex.export_filename(
            "max_temperature", "fleet", NOW,
            period_text="2026-08-08 to 2026-09-07 UTC",
        )
        assert name == (
            "max_temperature_fleet_2026-08-08_to_2026-09-07_UTC_"
            "20260825T120000Z.csv"
        )

    def test_pdf_filename_uses_the_pdf_extension(self):
        doc = ex.installed_rtls_document([], scope_label="fleet", now=NOW)
        _content, mime, filename = ex.render_export("installed_rtls", doc, "pdf")
        assert mime == "application/pdf"
        assert filename.endswith(".pdf")

    def test_xlsx_filename_and_mime_type(self):
        doc = ex.installed_rtls_document([], scope_label="fleet", now=NOW)
        content, mime, filename = ex.render_export("installed_rtls", doc, "xlsx")
        assert mime == (
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        )
        assert filename.endswith(".xlsx")
        assert isinstance(content, bytes)


@dataclass(frozen=True)
class FakeMaxTempRow:
    ou = None
    zone = None
    sector = None
    cnc = None
    feeder_name = None
    transformer: str = "tc1code"
    date_installed: datetime | None = datetime(2025, 1, 1, tzinfo=timezone.utc)
    max_temperature_at: datetime | None = datetime(2026, 9, 1, tzinfo=timezone.utc)
    max_temperature: float | None = 91.4


# ---------------------------------------------------------------------------
# R4-D11 — Maximum Temperature's period identification lives in its
# filename and its PDF header line, NEVER inside the CSV body, and the
# other two reports' CSV/filename behavior is completely unaffected
# ---------------------------------------------------------------------------


class TestMaxTemperaturePeriodIdentification:
    PERIOD = "2026-08-08 to 2026-09-07 UTC"

    def test_csv_first_row_is_exactly_the_column_contract_no_preamble(self):
        """No metadata preamble of any kind — Maximum Temperature's CSV is
        a plain rectangular table, identical in shape to the other two."""
        doc = ex.max_temperature_document(
            [FakeMaxTempRow()], scope_label="Entire Fleet",
            period_text=self.PERIOD, now=NOW,
        )
        parsed = _parse(ex.format_csv(doc))
        assert parsed[0] == list(get_report("max_temperature").columns)
        assert len(parsed) == 2                      # header + one data row
        assert parsed[1][5] == "tc1code"

    def test_csv_rows_all_have_exactly_the_contract_column_count(self):
        doc = ex.max_temperature_document(
            [FakeMaxTempRow()], scope_label="fleet",
            period_text=self.PERIOD, now=NOW,
        )
        parsed = _parse(ex.format_csv(doc))
        ncols = len(get_report("max_temperature").columns)
        assert all(len(row) == ncols for row in parsed)

    def test_period_text_never_appears_anywhere_in_the_csv_bytes(self):
        doc = ex.max_temperature_document(
            [FakeMaxTempRow()], scope_label="fleet",
            period_text=self.PERIOD, now=NOW,
        )
        assert self.PERIOD not in ex.format_csv(doc)

    def test_installed_rtls_and_rtl_alarms_csv_are_unaffected_by_r4d11(self):
        """Neither pre-existing report's CSV changes shape, even though
        both now carry a static `period_text` for the PDF's benefit."""
        installed_csv = ex.format_csv(
            ex.installed_rtls_document([], scope_label="fleet", now=NOW)
        )
        alarms_csv = ex.format_csv(
            ex.rtl_alarms_document([], scope_label="fleet", now=NOW)
        )
        assert _parse(installed_csv) == [list(get_report("installed_rtls").columns)]
        assert _parse(alarms_csv) == [list(get_report("rtl_alarms_30d").columns)]

    def test_installed_rtls_and_rtl_alarms_filenames_are_unchanged(self):
        """The period segment never reaches these two reports' filenames,
        matching test_shape_is_deterministic exactly."""
        doc1 = ex.installed_rtls_document([], scope_label="fleet", now=NOW)
        doc2 = ex.rtl_alarms_document([], scope_label="fleet", now=NOW)
        _c1, _m1, name1 = ex.render_export("installed_rtls", doc1)
        _c2, _m2, name2 = ex.render_export("rtl_alarms_30d", doc2)
        assert name1 == "installed_rtls_fleet_20260825T120000Z.csv"
        assert name2 == "rtl_alarms_30d_fleet_20260825T120000Z.csv"

    def test_max_temperature_filename_period_segment_is_deterministic(self):
        """Same inputs -> the same filename, every time — the period
        segment is derived purely from the resolved [since, until], not
        from wall-clock or randomness."""
        doc = ex.max_temperature_document(
            [], scope_label="fleet", period_text=self.PERIOD, now=NOW,
        )
        _c1, _m1, name1 = ex.render_export("max_temperature", doc)
        _c2, _m2, name2 = ex.render_export("max_temperature", doc)
        assert name1 == name2
        assert self.PERIOD.replace(" ", "_") in name1


# ---------------------------------------------------------------------------
# R4-D10 — the PDF formatter
# ---------------------------------------------------------------------------


def _spy_cell_texts(monkeypatch) -> list[str]:
    """Capture every text string passed to FPDF.cell (the header block:
    title/scope/period/generated), while still letting the real call
    render — behavior verification, not a fake PDF."""
    from fpdf import FPDF

    texts: list[str] = []
    original_cell = FPDF.cell

    def spy_cell(self, w=None, h=None, text="", *args, **kwargs):
        texts.append(text)
        return original_cell(self, w, h, text, *args, **kwargs)

    monkeypatch.setattr(FPDF, "cell", spy_cell)
    return texts


class TestPdfFormatter:
    def test_output_is_a_real_pdf(self):
        doc = ex.max_temperature_document(
            [FakeMaxTempRow()], scope_label="Entire Fleet",
            period_text="2026-08-08 to 2026-09-07 UTC", now=NOW,
        )
        content = ex.format_pdf(doc)
        assert isinstance(content, bytes)
        assert content.startswith(b"%PDF-")
        assert content.rstrip().endswith(b"%%EOF")

    # R4-D12: format_pdf uses fpdf2's normal (compressed) output, so its
    # content streams cannot be substring-matched as raw text. The tests
    # below verify FORMATTER BEHAVIOR instead — the exact header/cell
    # values actually handed to fpdf2's `Table.row`, via a spy that still
    # lets the real call through (rendering is unaffected; nothing here
    # is a fake PDF) — plus structural framing that survives compression
    # (`/Type /Page` object markers are never inside a compressed stream).

    def test_title_scope_period_and_generated_time_are_all_written(self, monkeypatch):
        """PDF continues to display Report/Scope/Period/Generated inside
        the document (preserved) — verified as formatter behavior (the
        exact text handed to FPDF.cell), not by grepping compressed
        bytes."""
        texts = _spy_cell_texts(monkeypatch)
        doc = ex.max_temperature_document(
            [], scope_label="Entire Fleet",
            period_text="2026-08-08 to 2026-09-07 UTC", now=NOW,
        )
        ex.format_pdf(doc)

        assert "Maximum Temperature" in texts             # title
        assert "Scope: Entire Fleet" in texts
        assert "Period: 2026-08-08 to 2026-09-07 UTC" in texts
        assert "Generated: 2026-08-25T12:00:00Z" in texts

    def test_report_with_no_period_omits_the_period_line_but_not_scope(self, monkeypatch):
        """Proves the omission path (`period_text is None`) independently,
        using a hand-built document — no real report currently has this
        shape (R4-D11 gives all three a period_text), but the formatter
        must still handle it correctly."""
        texts = _spy_cell_texts(monkeypatch)
        doc = ex.ExportDocument(
            report_key="installed_rtls", title="Installed RTLs",
            scope_label="Entire Fleet", generated_at_utc=NOW,
            headers=("UID",), rows=(), period_text=None,
        )
        ex.format_pdf(doc)

        assert "Scope: Entire Fleet" in texts
        assert not any(t.startswith("Period:") for t in texts)

    def test_table_receives_the_exact_headers_then_one_row_per_data_row(self, monkeypatch):
        from fpdf.table import Table

        calls = []
        original_row = Table.row

        def spy_row(self, cells=(), style=None):
            calls.append(list(cells))
            return original_row(self, cells, style)

        monkeypatch.setattr(Table, "row", spy_row)

        doc = ex.max_temperature_document(
            [FakeMaxTempRow(), FakeMaxTempRow(transformer="tc2")],
            scope_label="fleet",
            period_text="2026-08-08 to 2026-09-07 UTC", now=NOW,
        )
        content = ex.format_pdf(doc)

        assert content.startswith(b"%PDF-")            # real render happened
        assert len(calls) == 1 + len(doc.rows)          # header + N data rows
        assert calls[0] == list(doc.headers)
        assert calls[1] == [ex._cell(v) for v in doc.rows[0]]
        assert calls[2] == [ex._cell(v) for v in doc.rows[1]]

    def test_none_cells_reach_the_table_as_empty_strings(self, monkeypatch):
        from fpdf.table import Table

        calls = []
        monkeypatch.setattr(
            Table, "row",
            lambda self, cells=(), style=None: calls.append(list(cells)),
        )

        row = FakeMaxTempRow(date_installed=None, max_temperature_at=None,
                              max_temperature=None)
        doc = ex.max_temperature_document(
            [row], scope_label="fleet",
            period_text="2026-08-08 to 2026-09-07 UTC", now=NOW,
        )
        ex.format_pdf(doc)

        assert calls[1] == ["", "", "", "", "", "tc1code", "", "", ""]

    def test_zero_rows_sends_only_the_header_row_to_the_table(self, monkeypatch):
        from fpdf.table import Table

        calls = []
        original_row = Table.row
        monkeypatch.setattr(
            Table, "row",
            lambda self, cells=(), style=None: (
                calls.append(list(cells)), original_row(self, cells, style)
            )[1],
        )

        doc = ex.rtl_alarms_document([], scope_label="fleet", now=NOW)
        content = ex.format_pdf(doc)

        assert content.startswith(b"%PDF-")
        assert calls == [list(doc.headers)]

    def test_none_cells_render_without_error(self):
        row = FakeMaxTempRow(date_installed=None, max_temperature_at=None,
                              max_temperature=None)
        doc = ex.max_temperature_document(
            [row], scope_label="fleet",
            period_text="2026-08-08 to 2026-09-07 UTC", now=NOW,
        )
        content = ex.format_pdf(doc)   # must not raise
        assert content.startswith(b"%PDF-")

    def test_zero_rows_still_produces_a_valid_pdf(self):
        """R4-D5's convention extends to PDF: a legitimate empty report,
        never an exception."""
        doc = ex.rtl_alarms_document([], scope_label="fleet", now=NOW)
        content = ex.format_pdf(doc)
        assert content.startswith(b"%PDF-")
        assert content.count(b"/Type /Page") >= 1

    def test_many_rows_paginate_with_repeated_headers(self):
        """Real pagination, not a single unbounded page: more page
        objects exist for a document that cannot fit on one page than for
        one that fits easily."""
        doc_small = ex.max_temperature_document(
            [FakeMaxTempRow()], scope_label="fleet",
            period_text="p", now=NOW,
        )
        doc_big = ex.max_temperature_document(
            [FakeMaxTempRow(transformer=f"tc{i}") for i in range(400)],
            scope_label="fleet", period_text="p", now=NOW,
        )
        small_pages = ex.format_pdf(doc_small).count(b"/Type /Page")
        big_pages = ex.format_pdf(doc_big).count(b"/Type /Page")
        assert big_pages > small_pages
        assert big_pages > 2   # more than a trivial 1-page + tree count


# ---------------------------------------------------------------------------
# R4-D13/R4-D14 — the native XLSX formatter (FS-EXPORT-1)
# ---------------------------------------------------------------------------


class TestWorksheetName:
    def test_every_real_report_key_produces_a_valid_sheet_name(self):
        for key in ex.EXPORTABLE_REPORTS:
            name = ex._worksheet_name(key)
            assert 0 < len(name) <= 31
            assert not set(name) & ex._INVALID_SHEET_NAME_CHARS

    def test_hostile_report_key_is_sanitized_and_bounded(self):
        name = ex._worksheet_name("a" * 50 + "[bad]:name/here?")
        assert len(name) <= 31
        assert not set(name) & ex._INVALID_SHEET_NAME_CHARS


class TestXlsxFormatter:
    def _load(self, content: bytes):
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(content))

    def test_output_is_bytes_and_a_real_openable_workbook(self):
        doc = ex.max_temperature_document(
            [FakeMaxTempRow()], scope_label="Entire Fleet",
            period_text="2026-08-08 to 2026-09-07 UTC", now=NOW,
        )
        content = ex.format_xlsx(doc)
        assert isinstance(content, bytes)
        wb = self._load(content)
        assert wb.sheetnames == ["max_temperature"]

    @pytest.mark.parametrize(
        "report_key,builder",
        [
            ("installed_rtls", ex.installed_rtls_document),
            ("rtl_alarms_30d", ex.rtl_alarms_document),
        ],
    )
    def test_each_report_exports_a_valid_workbook_with_the_expected_sheet(
        self, report_key, builder,
    ):
        doc = builder([], scope_label="fleet", now=NOW)
        wb = self._load(ex.format_xlsx(doc))
        sheet = wb.active
        assert sheet.title == report_key
        assert [c.value for c in sheet[1]] == list(get_report(report_key).columns)

    def test_headers_match_the_client_contract_exactly_in_order(self):
        """R4-D14/task item 4: the exact config/reports.py column order,
        the same single source of truth CSV/PDF already use."""
        doc = ex.rtl_alarms_document([], scope_label="fleet", now=NOW)
        wb = self._load(ex.format_xlsx(doc))
        sheet = wb.active
        assert [c.value for c in sheet[1]] == list(
            get_report("rtl_alarms_30d").columns
        )

    def test_headers_are_bold(self):
        doc = ex.rtl_alarms_document([], scope_label="fleet", now=NOW)
        wb = self._load(ex.format_xlsx(doc))
        sheet = wb.active
        assert all(cell.font.bold for cell in sheet[1])

    def test_zero_rows_still_produces_a_valid_header_only_workbook(self):
        """R4-D5's convention extends to XLSX: a legitimate empty report,
        never an exception, and still directly openable."""
        doc = ex.rtl_alarms_document([], scope_label="fleet", now=NOW)
        wb = self._load(ex.format_xlsx(doc))
        sheet = wb.active
        assert sheet.max_row == 1
        assert [c.value for c in sheet[1]] == list(
            get_report("rtl_alarms_30d").columns
        )

    def test_no_metadata_preamble_row_one_is_the_header(self):
        """R4-D14: mirrors CSV's shape (R4-D11) — the table starts at A1,
        no title/scope/period block above it, unlike PDF."""
        doc = ex.max_temperature_document(
            [FakeMaxTempRow()], scope_label="Entire Fleet",
            period_text="2026-08-08 to 2026-09-07 UTC", now=NOW,
        )
        wb = self._load(ex.format_xlsx(doc))
        sheet = wb.active
        assert sheet["A1"].value == "OU"
        assert sheet.max_row == 2   # header + exactly one data row

    def test_row_count_matches_document_row_count(self):
        doc = ex.rtl_alarms_document(
            [FakeAlarmRow(), FakeAlarmRow(transformer="tc2")],
            scope_label="fleet", now=NOW,
        )
        wb = self._load(ex.format_xlsx(doc))
        assert wb.active.max_row == 1 + 2

    def test_string_cells_are_real_text_not_decorated(self):
        doc = ex.rtl_alarms_document([FakeAlarmRow()], scope_label="fleet", now=NOW)
        sheet = self._load(ex.format_xlsx(doc)).active
        cell = sheet.cell(row=2, column=6)   # Transformer
        assert cell.value == "tc1"
        assert isinstance(cell.value, str)

    def test_numeric_cells_are_real_numbers_not_decorated_strings(self):
        """Task item 5: numeric values -> numeric cells, never str()'d."""
        doc = ex.rtl_alarms_document([FakeAlarmRow()], scope_label="fleet", now=NOW)
        sheet = self._load(ex.format_xlsx(doc)).active
        battery_cell = sheet.cell(row=2, column=8)     # Battery(V)
        temperature_cell = sheet.cell(row=2, column=10)  # Temperature (C)
        assert battery_cell.value == pytest.approx(3.52)
        assert isinstance(battery_cell.value, (int, float))
        assert temperature_cell.value == pytest.approx(85.25)
        assert isinstance(temperature_cell.value, (int, float))

    def test_datetime_cells_are_real_excel_datetimes_labelled_utc(self):
        """Task item 5: a real Excel datetime, not an ISO-8601 string —
        the cell's own number format documents the UTC representation."""
        doc = ex.rtl_alarms_document([FakeAlarmRow()], scope_label="fleet", now=NOW)
        sheet = self._load(ex.format_xlsx(doc)).active
        cell = sheet.cell(row=2, column=9)   # Alarm Date & Time
        assert isinstance(cell.value, datetime)
        assert cell.value == datetime(2026, 8, 20, 6, 30)
        assert cell.value.tzinfo is None    # Excel datetimes carry no tz
        assert "UTC" in cell.number_format

    def test_naive_datetime_is_interpreted_as_utc(self):
        naive = FakeAlarmRow(alarm_at=datetime(2026, 8, 20, 6, 30))
        doc = ex.rtl_alarms_document([naive], scope_label="fleet", now=NOW)
        sheet = self._load(ex.format_xlsx(doc)).active
        assert sheet.cell(row=2, column=9).value == datetime(2026, 8, 20, 6, 30)

    def test_none_cells_are_truly_blank_not_empty_string(self):
        """Task item 5: None -> a blank cell, never a decorated/empty-
        string placeholder — a real Excel COUNTA()/SUM() must see nothing."""
        row = FakeAlarmRow(battery_voltage=None, temperature=None,
                            firmware_version=None, uid=None)
        doc = ex.rtl_alarms_document([row], scope_label="fleet", now=NOW)
        sheet = self._load(ex.format_xlsx(doc)).active
        for col in (7, 8, 10, 12):   # UID, Battery(V), Temperature, Firmware
            assert sheet.cell(row=2, column=col).value is None

    def test_taxonomy_cells_are_blank(self):
        doc = ex.rtl_alarms_document([FakeAlarmRow()], scope_label="fleet", now=NOW)
        sheet = self._load(ex.format_xlsx(doc)).active
        for col in range(1, 6):   # OU, Zone, Sector, CNC, Feeder
            assert sheet.cell(row=2, column=col).value is None

    def test_max_temperature_representative_row_round_trips(self):
        doc = ex.max_temperature_document(
            [FakeMaxTempRow()], scope_label="fleet",
            period_text="2026-08-08 to 2026-09-07 UTC", now=NOW,
        )
        sheet = self._load(ex.format_xlsx(doc)).active
        assert sheet.cell(row=2, column=6).value == "tc1code"     # Transformer
        assert sheet.cell(row=2, column=7).value == datetime(2025, 1, 1)
        assert sheet.cell(row=2, column=8).value == datetime(2026, 9, 1)
        assert sheet.cell(row=2, column=9).value == pytest.approx(91.4)

    def test_period_text_never_appears_in_the_workbook_headers(self):
        """R4-D14: like CSV (R4-D11), XLSX carries no period preamble —
        only the header row exists above the data."""
        period = "2026-08-08 to 2026-09-07 UTC"
        doc = ex.max_temperature_document(
            [FakeMaxTempRow()], scope_label="fleet", period_text=period, now=NOW,
        )
        sheet = self._load(ex.format_xlsx(doc)).active
        header_and_first_row = [c.value for c in sheet[1]] + [
            c.value for c in sheet[2]
        ]
        assert period not in [str(v) for v in header_and_first_row]
